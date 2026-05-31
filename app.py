from flask import Flask, render_template, request
from sqlalchemy import create_engine, text
import pandas as pd
from flask import request, jsonify
import os
from datetime import datetime
from flask import send_file
from openpyxl import load_workbook
import shutil


app = Flask(__name__)

DB_USER = "postgres"
DB_PASSWORD = "faraz9024"
DB_HOST = "localhost"
DB_PORT = "5432"
DB_NAME = "epl_portal"

engine = create_engine(
    f"postgresql+psycopg2://{DB_USER}:{DB_PASSWORD}@{DB_HOST}:{DB_PORT}/{DB_NAME}"
)

@app.route("/site")
def site_search():
    site_id = request.args.get("site_id")

    if not site_id:
        return render_template("site_search.html")

    site = pd.read_sql(
        f"""
        SELECT *
        FROM site_master
        WHERE UPPER(site_id)=UPPER('{site_id}')
        """,
        engine
    )

    lte = pd.read_sql(
        f"""
        SELECT *
        FROM lte_cells
        WHERE UPPER(site_id)=UPPER('{site_id}')
        """,
        engine
    )

    l2600 = pd.read_sql(
        f"""
        SELECT *
        FROM l2600_cells
        WHERE UPPER(site_id)=UPPER('{site_id}')
        """,
        engine
    )

    nr = pd.read_sql(
        f"""
        SELECT *
        FROM nr_cells
        WHERE UPPER(site_id)=UPPER('{site_id}')
        """,
        engine
    )

    return render_template(
        "site_search.html",
        site=site,
        lte=lte,
        l2600=l2600,
        nr=nr,
        site_id=site_id
    )

@app.route("/cell")
def cell_search():
    cell_name = request.args.get("cell_name")

    if not cell_name:
        return render_template("cell_search.html")

    # LTE Search
    cell = pd.read_sql(
        f"""
        SELECT *
        FROM lte_cells
        WHERE UPPER(cell_name)=UPPER('{cell_name}')
        """,
        engine
    )

    if not cell.empty:
        return render_template(
            "cell_search.html",
            cell=cell,
            cell_name=cell_name,
            tech="LTE"
        )

    # L2600 Search
    cell = pd.read_sql(
        f"""
        SELECT *
        FROM l2600_cells
        WHERE UPPER(cell_name)=UPPER('{cell_name}')
        """,
        engine
    )

    if not cell.empty:
        return render_template(
            "cell_search.html",
            cell=cell,
            cell_name=cell_name,
            tech="L2600"
        )

    # NR Search
    cell = pd.read_sql(
        f"""
        SELECT *
        FROM nr_cells
        WHERE UPPER(cell_name)=UPPER('{cell_name}')
        """,
        engine
    )

    if not cell.empty:
        return render_template(
            "cell_search.html",
            cell=cell,
            cell_name=cell_name,
            tech="NR"
        )

    return render_template(
        "cell_search.html",
        cell_name=cell_name
    )
    
@app.route("/")
def dashboard():
    total_sites = pd.read_sql("SELECT COUNT(*) cnt FROM site_master", engine)["cnt"][0]
    lte_sites = pd.read_sql("SELECT COUNT(DISTINCT site_id) cnt FROM lte_cells", engine)["cnt"][0]
    l2600_sites = pd.read_sql("SELECT COUNT(DISTINCT site_id) cnt FROM l2600_cells", engine)["cnt"][0]
    nr_sites = pd.read_sql("SELECT COUNT(DISTINCT site_id) cnt FROM nr_cells", engine)["cnt"][0]
    lte_cells = pd.read_sql("SELECT COUNT(*) cnt FROM lte_cells", engine)["cnt"][0]
    l2600_cells = pd.read_sql("SELECT COUNT(*) cnt FROM l2600_cells", engine)["cnt"][0]
    nr_cells = pd.read_sql("SELECT COUNT(*) cnt FROM nr_cells", engine)["cnt"][0]
    lte_on = pd.read_sql("SELECT COUNT(*) cnt FROM lte_cells WHERE onair = 1", engine)["cnt"][0]
    lte_off = pd.read_sql("SELECT COUNT(*) cnt FROM lte_cells WHERE onair = 0", engine)["cnt"][0]
    l2600_on = pd.read_sql("SELECT COUNT(*) cnt FROM l2600_cells WHERE onair = 1", engine)["cnt"][0]
    l2600_off = pd.read_sql("SELECT COUNT(*) cnt FROM l2600_cells WHERE onair = 0", engine)["cnt"][0]
    nr_on = pd.read_sql("SELECT COUNT(*) cnt FROM nr_cells WHERE onair = 1", engine)["cnt"][0]
    nr_off = pd.read_sql("SELECT COUNT(*) cnt FROM nr_cells WHERE onair = 0", engine)["cnt"][0]
    
    current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    return render_template(
        "dashboard.html",
        total_sites=total_sites,
        lte_sites=lte_sites,
        l2600_sites=l2600_sites,
        nr_sites=nr_sites,
        lte_cells=lte_cells,
        l2600_cells=l2600_cells,
        nr_cells=nr_cells,
        lte_on=lte_on,
        lte_off=lte_off,
        l2600_on=l2600_on,
        l2600_off=l2600_off,
        nr_on=nr_on,
        nr_off=nr_off,
        current_time=current_time
    )

# ===========================
# SITE CELLS API
# ===========================

@app.route("/api/site_cells/<site_id>")
def site_cells(site_id):
    lte = pd.read_sql(
        f"""
        SELECT cell_name, azimuth, mtilt
        FROM lte_cells
        WHERE UPPER(site_id)=UPPER('{site_id}')
        """,
        engine
    )

    l2600 = pd.read_sql(
        f"""
        SELECT cell_name, azimuth, mtilt
        FROM l2600_cells
        WHERE UPPER(site_id)=UPPER('{site_id}')
        """,
        engine
    )

    nr = pd.read_sql(
        f"""
        SELECT cell_name, azimuth, mtilt
        FROM nr_cells
        WHERE UPPER(site_id)=UPPER('{site_id}')
        """,
        engine
    )

    all_cells = pd.concat([lte, l2600, nr])
    all_cells = all_cells.sort_values("cell_name")

    return {"cells": all_cells.to_dict(orient="records")}

@app.route("/api/update_cells", methods=["POST"])
def update_cells():
    data = request.json
    cells = data.get("cells", [])
    engineer = data.get("engineer_name")
    rigger = data.get("rigger_name")
    site_id = data.get("site_id")

    conn = engine.connect()

    for cell in cells:
        cell_name = cell["cell_name"]
        new_azimuth = cell["azimuth"]
        new_mtilt = cell["mtilt"]

        # LTE
        old = pd.read_sql(
            f"""
            SELECT azimuth,mtilt
            FROM lte_cells
            WHERE cell_name='{cell_name}'
            """,
            engine
        )

        if not old.empty:
            old_az = old.iloc[0]["azimuth"]
            old_tilt = old.iloc[0]["mtilt"]

            conn.execute(text(f"""
                UPDATE lte_cells
                SET azimuth={new_azimuth},
                    mtilt='{new_mtilt}'
                WHERE cell_name='{cell_name}'
            """))

        else:
            # L2600
            old = pd.read_sql(
                f"""
                SELECT azimuth,mtilt
                FROM l2600_cells
                WHERE cell_name='{cell_name}'
                """,
                engine
            )

            if not old.empty:
                old_az = old.iloc[0]["azimuth"]
                old_tilt = old.iloc[0]["mtilt"]

                conn.execute(text(f"""
                    UPDATE l2600_cells
                    SET azimuth={new_azimuth},
                        mtilt='{new_mtilt}'
                    WHERE cell_name='{cell_name}'
                """))

            else:
                # NR
                old = pd.read_sql(
                    f"""
                    SELECT azimuth,mtilt
                    FROM nr_cells
                    WHERE cell_name='{cell_name}'
                    """,
                    engine
                )

                if not old.empty:
                    old_az = old.iloc[0]["azimuth"]
                    old_tilt = old.iloc[0]["mtilt"]

                    conn.execute(text(f"""
                        UPDATE nr_cells
                        SET azimuth={new_azimuth},
                            mtilt='{new_mtilt}'
                        WHERE cell_name='{cell_name}'
                    """))

        # Insert into physical_change_history table (your existing table)
        conn.execute(text(f"""
            INSERT INTO physical_change_history
            (site_id, cell_name, old_azimuth, old_mtilt, new_azimuth, new_mtilt, engineer_name, rigger_name)
            VALUES
            ('{site_id}', '{cell_name}', {old_az}, '{old_tilt}', {new_azimuth}, '{new_mtilt}', '{engineer}', '{rigger}')
        """))

    conn.commit()
    return jsonify({"status": "success"})

@app.route('/download/epl')
def download_epl():
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    template_file = r"D:\EPL_Portal\data\EPL.xlsx"
    output_file = rf"D:\EPL_Portal\downloads\EPL_Updated_{timestamp}.xlsx"
    
    shutil.copy(template_file, output_file)
    wb = load_workbook(output_file)
    
    # FDD CELLS (LTE)
    ws_fdd = wb["FDD Cells"]
    lte_df = pd.read_sql("SELECT cell_name, azimuth, mtilt FROM lte_cells", engine)
    lte_lookup = {str(row["cell_name"]).strip().upper(): {"azimuth": row["azimuth"], "mtilt": row["mtilt"]} 
                  for _, row in lte_df.iterrows()}
    
    for row in range(2, ws_fdd.max_row + 1):
        cell_name = ws_fdd.cell(row=row, column=1).value
        if cell_name:
            key = str(cell_name).strip().upper()
            if key in lte_lookup:
                rec = lte_lookup[key]
                ws_fdd.cell(row=row, column=25).value = rec["azimuth"]
                ws_fdd.cell(row=row, column=31).value = rec["mtilt"]
    
    # TDD 2.6 (L2600)
    ws_tdd = wb["TDD 2.6"]
    tdd_df = pd.read_sql("SELECT cell_name, azimuth, mtilt FROM l2600_cells", engine)
    tdd_lookup = {str(row["cell_name"]).strip().upper(): {"azimuth": row["azimuth"], "mtilt": row["mtilt"]} 
                  for _, row in tdd_df.iterrows()}
    
    for row in range(2, ws_tdd.max_row + 1):
        cell_name = ws_tdd.cell(row=row, column=1).value
        if cell_name:
            key = str(cell_name).strip().upper()
            if key in tdd_lookup:
                rec = tdd_lookup[key]
                ws_tdd.cell(row=row, column=25).value = rec["azimuth"]
                ws_tdd.cell(row=row, column=31).value = rec["mtilt"]
    
    # NR CELL
    ws_nr = wb["NR Cell"]
    nr_df = pd.read_sql("SELECT cell_name, azimuth, mtilt FROM nr_cells", engine)
    nr_lookup = {str(row["cell_name"]).strip().upper(): {"azimuth": row["azimuth"], "mtilt": row["mtilt"]} 
                 for _, row in nr_df.iterrows()}
    
    for row in range(2, ws_nr.max_row + 1):
        cell_name = ws_nr.cell(row=row, column=11).value
        if cell_name:
            key = str(cell_name).strip().upper()
            if key in nr_lookup:
                rec = nr_lookup[key]
                ws_nr.cell(row=row, column=13).value = rec["azimuth"]
                ws_nr.cell(row=row, column=19).value = rec["mtilt"]
    
    wb.save(output_file)
    return send_file(output_file, as_attachment=True, download_name=f"EPL_Updated_{timestamp}.xlsx")
@app.route("/history")
def change_history():
    """View change history from physical_change_history table"""
    
    # Get filter parameters
    site_filter = request.args.get("site_id", "")
    cell_filter = request.args.get("cell_name", "")
    engineer_filter = request.args.get("engineer", "")
    
    # Build the query - using approved_at as timestamp
    query = """
        SELECT 
            id,
            site_id,
            cell_name,
            old_azimuth,
            old_mtilt,
            new_azimuth,
            new_mtilt,
            engineer_name,
            rigger_name,
            approved_at
        FROM physical_change_history
        WHERE 1=1
    """
    params = {}
    
    if site_filter:
        query += " AND UPPER(site_id) = UPPER(:site_id)"
        params["site_id"] = site_filter
    
    if cell_filter:
        query += " AND UPPER(cell_name) = UPPER(:cell_name)"
        params["cell_name"] = cell_filter
    
    if engineer_filter:
        query += " AND UPPER(engineer_name) = UPPER(:engineer_name)"
        params["engineer"] = engineer_filter
    
    query += " ORDER BY approved_at DESC LIMIT 500"
    
    # Fetch history data
    history_df = pd.read_sql(text(query), engine, params=params)
    
    # Get unique values for filters
    sites = pd.read_sql("SELECT DISTINCT site_id FROM physical_change_history WHERE site_id IS NOT NULL ORDER BY site_id", engine)
    engineers = pd.read_sql("SELECT DISTINCT engineer_name FROM physical_change_history WHERE engineer_name IS NOT NULL ORDER BY engineer_name", engine)
    
    return render_template(
        "change_history.html",
        history=history_df,
        sites=sites['site_id'].tolist() if not sites.empty else [],
        engineers=engineers['engineer_name'].tolist() if not engineers.empty else [],
        site_filter=site_filter,
        cell_filter=cell_filter,
        engineer_filter=engineer_filter
    )
@app.route("/download/history")
def download_history():
    """Download change history as Excel"""
    
    # Get filter parameters
    site_filter = request.args.get("site_id", "")
    cell_filter = request.args.get("cell_name", "")
    engineer_filter = request.args.get("engineer", "")
    
    # Build query
    query = """
        SELECT 
            id,
            site_id,
            cell_name,
            old_azimuth,
            old_mtilt,
            new_azimuth,
            new_mtilt,
            engineer_name,
            rigger_name,
            approved_at
        FROM physical_change_history
        WHERE 1=1
    """
    params = {}
    
    if site_filter:
        query += " AND UPPER(site_id) = UPPER(:site_id)"
        params["site_id"] = site_filter
    
    if cell_filter:
        query += " AND UPPER(cell_name) = UPPER(:cell_name)"
        params["cell_name"] = cell_filter
    
    if engineer_filter:
        query += " AND UPPER(engineer_name) = UPPER(:engineer_name)"
        params["engineer"] = engineer_filter
    
    query += " ORDER BY approved_at DESC"
    
    history_df = pd.read_sql(text(query), engine, params=params)
    
    # Create Excel file
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = rf"D:\EPL_Portal\downloads\Change_History_{timestamp}.xlsx"
    
    # Ensure directory exists
    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        history_df.to_excel(writer, sheet_name='Change History', index=False)
        
        # Auto-adjust column widths
        worksheet = writer.sheets['Change History']
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
    
    return send_file(
        output_file,
        as_attachment=True,
        download_name=f"Change_History_{timestamp}.xlsx"
    )
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)